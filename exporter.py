"""
exporter.py — Export ReviewResult as a formatted .docx peer review document.
Matches real journal review format (Summary / Strengths / Major / Minor / Verdict).
"""

from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reviewer import ReviewResult

_SEVERITY_COLORS = {
    "critical": RGBColor(0xA3, 0x2D, 0x2D),
    "major":    RGBColor(0xBA, 0x75, 0x17),
    "minor":    RGBColor(0x3B, 0x6D, 0x11),
}

_VERDICT_LABELS = {
    "accept":           "ACCEPT",
    "minor_revisions":  "ACCEPT WITH MINOR REVISIONS",
    "major_revisions":  "MAJOR REVISIONS REQUIRED",
    "reject":           "REJECT",
}


def _heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.color.rgb = RGBColor(0x1C, 0x1C, 0x1C)


def _body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)


def _issue_block(doc: Document, label: str, quote: str, explanation: str,
                 suggestion: str, severity: str):
    p = doc.add_paragraph()
    run = p.add_run(f"[{severity.upper()}] {label}  ")
    run.bold = True
    run.font.color.rgb = _SEVERITY_COLORS.get(severity, RGBColor(0, 0, 0))

    if quote:
        q = doc.add_paragraph(f'"{quote}"')
        q.paragraph_format.left_indent = Inches(0.3)
        q.runs[0].italic = True
        q.paragraph_format.space_after = Pt(2)

    exp = doc.add_paragraph(explanation)
    exp.paragraph_format.left_indent = Inches(0.3)
    exp.paragraph_format.space_after = Pt(2)

    sug = doc.add_paragraph(f"Suggestion: {suggestion}")
    sug.paragraph_format.left_indent = Inches(0.3)
    sug.runs[0].font.color.rgb = RGBColor(0x0C, 0x44, 0x7C)
    sug.paragraph_format.space_after = Pt(8)


def export_review_docx(result: ReviewResult, doc_name: str = "document") -> bytes:
    doc = Document()

    # Margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)

    # Title block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PEER REVIEW REPORT")
    run.bold = True
    run.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"Document: {doc_name}").font.size = Pt(11)
    doc.add_paragraph()

    # Verdict banner
    verdict_label = _VERDICT_LABELS.get(result.verdict, result.verdict.upper())
    vp = doc.add_paragraph()
    vrun = vp.add_run(f"RECOMMENDATION: {verdict_label}   |   Score: {result.overall_score}/10")
    vrun.bold = True
    vrun.font.size = Pt(12)
    color = {
        "accept": RGBColor(0x0B, 0x80, 0x4B),
        "minor_revisions": RGBColor(0x63, 0x38, 0x06),
        "major_revisions": RGBColor(0xBA, 0x75, 0x17),
        "reject": RGBColor(0xA3, 0x2D, 0x2D),
    }.get(result.verdict, RGBColor(0, 0, 0))
    vrun.font.color.rgb = color
    doc.add_paragraph()

    # Document map
    _heading(doc, "Document Overview", level=1)
    _body(doc, f"Thesis: {result.doc_map.thesis}")
    _body(doc, f"Type: {result.doc_map.doc_type}  |  Methodology: {result.doc_map.methodology}")
    _body(doc, f"Stated contribution: {result.doc_map.contribution}")
    doc.add_paragraph()

    # Summary
    _heading(doc, "Summary", level=1)
    _body(doc, result.summary)
    doc.add_paragraph()

    # Strengths
    _heading(doc, "Strengths", level=1)
    for s in result.strengths:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_paragraph()

    # Major concerns — claim issues
    criticals = [i for i in result.claim_issues if i.severity == "critical"]
    majors    = [i for i in result.claim_issues if i.severity == "major"]
    minors    = [i for i in result.claim_issues if i.severity == "minor"]

    _heading(doc, f"Major Concerns  ({len(criticals) + len(majors)} issues)", level=1)
    for concern in result.major_concerns:
        doc.add_paragraph(concern, style="List Bullet")
    doc.add_paragraph()

    _heading(doc, "Detailed Claim Issues", level=2)
    for issue in criticals + majors:
        _issue_block(
            doc,
            label=f"{issue.section} — {issue.issue_type.replace('_', ' ').title()}",
            quote=issue.quote,
            explanation=issue.explanation,
            suggestion=issue.suggestion,
            severity=issue.severity,
        )

    # Contradictions
    if result.contradictions:
        _heading(doc, f"Internal Contradictions  ({len(result.contradictions)} found)", level=1)
        for c in result.contradictions:
            p = doc.add_paragraph()
            run = p.add_run(f"[{c.severity.upper()}] {c.section_a}  ↔  {c.section_b}  ")
            run.bold = True
            run.font.color.rgb = _SEVERITY_COLORS.get(c.severity, RGBColor(0, 0, 0))

            q1 = doc.add_paragraph(f'Section A: "{c.quote_a}"')
            q1.paragraph_format.left_indent = Inches(0.3)
            q1.runs[0].italic = True

            q2 = doc.add_paragraph(f'Section B: "{c.quote_b}"')
            q2.paragraph_format.left_indent = Inches(0.3)
            q2.runs[0].italic = True

            exp = doc.add_paragraph(c.explanation)
            exp.paragraph_format.left_indent = Inches(0.3)
            exp.paragraph_format.space_after = Pt(8)
        doc.add_paragraph()

    # Minor concerns
    _heading(doc, f"Minor Concerns  ({len(minors)} issues)", level=1)
    for concern in result.minor_concerns:
        doc.add_paragraph(concern, style="List Bullet")
    doc.add_paragraph()
    for issue in minors:
        _issue_block(
            doc,
            label=f"{issue.section} — {issue.issue_type.replace('_', ' ').title()}",
            quote=issue.quote,
            explanation=issue.explanation,
            suggestion=issue.suggestion,
            severity=issue.severity,
        )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
